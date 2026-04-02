import os
from flask import Flask, request, jsonify
from werkzeug.utils import secure_filename
from pptx import Presentation
from PyPDF2 import PdfReader
from docx import Document
from transformers import pipeline

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = './temp_summaries'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# ------------------ 1. LOAD MODEL ONCE ------------------
# We load this outside the route so it stays in memory
print("Loading summarization model... please wait.")
summarizer = pipeline(
    "summarization", 
    model="sshleifer/distilbart-cnn-12-6",
    device=-1 # Set to 0 if you have a GPU
)

# ------------------ 2. EXTRACTION LOGIC ------------------

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == '.pdf':
            reader = PdfReader(file_path)
            return "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
        elif ext == '.docx':
            return "\n".join([para.text for para in Document(file_path).paragraphs])
        elif ext == '.pptx':
            prs = Presentation(file_path)
            return "\n".join([shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text")])
        return None
    except Exception as e:
        print(f"Error: {e}")
        return None

def chunk_text(text, max_words=300): # Increased slightly for better context
    words = text.split()
    return [" ".join(words[i:i+max_words]) for i in range(0, len(words), max_words)]

# ------------------ 3. THE ENDPOINT ------------------

@app.route('/summarize', methods=['POST'])
def handle_summarization():
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "Empty filename"}), 400

    # Save and Extract
    filename = secure_filename(file.filename)
    file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(file_path)
    
    full_text = extract_text(file_path)
    os.remove(file_path) # Clean up immediately

    if not full_text or len(full_text.strip()) < 10:
        return jsonify({"error": "File contains no readable text"}), 400

    # Process Summarization
    chunks = chunk_text(full_text)
    
    try:
        # Batch processing for speed
        raw_summaries = summarizer(
            chunks, 
            max_length=100, 
            min_length=30, 
            do_sample=False,
            batch_size=2
        )
        
        final_summary = " ".join([s['summary_text'] for s in raw_summaries])
        
        return jsonify({
            "filename": filename,
            "summary": final_summary,
            "original_chunks": len(chunks)
        })

    except Exception as e:
        return jsonify({"error": f"Summarization failed: {str(e)}"}), 500

if __name__ == '__main__':
    
    app.run(port=5002, debug=False)